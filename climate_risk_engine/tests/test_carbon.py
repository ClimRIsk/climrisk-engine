"""GHG inventory (carbon accounting) module tests."""
import io
import os

import pytest

os.environ.setdefault("CRI_API_KEYS", "")
from fastapi.testclient import TestClient  # noqa: E402

from cri.api.main import app  # noqa: E402
from cri.carbon import factors as F  # noqa: E402
from cri.carbon.assurance import assurance_workbook, parse_workbook, template_workbook  # noqa: E402
from cri.carbon.inventory import InventoryError, base_year_recalculation, calculate, from_dict  # noqa: E402

c = TestClient(app)
GAS = "desnz26:fuels-gaseous-fuels-natural-gas-kwh-gross-cv"
R410A = "desnz26:refrigerant-other-blends-r410a-emissions-including-only-kyoto-products-kg"
ALL_EXCLUDED = {f"s3_c{k}": {"status": "excluded", "reason": "test"} for k in range(1, 16)}


def inv(records, entities=None, **kw):
    d = {"organisation": "T", "reporting_year": 2025, "entities": entities or [{"id": "E", "name": "E", "country": "GBR"}],
         "records": records, "category_status": dict(ALL_EXCLUDED)}
    d.update(kw)
    return from_dict(d)


# ── factors ──────────────────────────────────────────────────────────────────

def test_desnz_factor_and_ar6_restatement():
    ar5 = F.desnz(GAS, "AR5")
    ar6 = F.desnz(GAS, "AR6")
    assert ar5["kgco2e_per_unit"] == pytest.approx(0.18231, rel=1e-3)
    assert ar6["kgco2e_per_unit"] > ar5["kgco2e_per_unit"] * 0.999        # CH4/N2O slightly higher on AR6
    assert ar6["kgco2e_per_unit"] < ar5["kgco2e_per_unit"] * 1.01


def test_refrigerant_ar6():
    assert F.desnz(R410A, "AR5")["kgco2e_per_unit"] == pytest.approx(1924, rel=1e-3)
    assert F.desnz(R410A, "AR6")["kgco2e_per_unit"] == pytest.approx(0.5 * 771 + 0.5 * 3740)


def test_grid_factors_by_region():
    assert F.electricity_location("GBR")["source"].startswith("UK DESNZ")
    tx = F.electricity_location("USA", us_state="TX")
    assert tx["subregion"] == "ERCT" and 250 < tx["g_per_kwh"] < 450
    assert F.electricity_location("FRA")["g_per_kwh"] < 60
    assert F.electricity_location("DEU")["g_per_kwh"] > 200
    assert "Ember" in F.electricity_location("IND")["source"]
    assert F.electricity_residual("DEU")["g_per_kwh"] > F.electricity_location("DEU")["g_per_kwh"]
    assert F.electricity_residual("USA", "ERCT") is not None
    assert F.electricity_residual("IND") is None


def test_spend_deflation():
    s = F.spend_factor("331110", 2022)
    assert s["deflator_to_2022"] == pytest.approx(1.0)
    assert F.spend_factor("331110", 2025)["kgco2e_per_usd_nominal"] < s["kgco2e_per_usd_nominal"]


# ── calculation rules ────────────────────────────────────────────────────────

def test_scope1_arithmetic_and_wtt_to_cat3():
    r = calculate(inv([{"entity": "E", "category": "s1_stationary", "method": "factor", "factor_id": GAS,
                        "quantity": 1_000_000, "evidence": "x"}], gwp="AR5"))
    assert r["totals_t"]["scope1"] == pytest.approx(182.31, rel=1e-3)
    assert r["scope3_by_category_t"]["Fuel- and energy-related activities"] == pytest.approx(30.21, rel=1e-2)


def test_unit_mismatch_rejected():
    with pytest.raises(InventoryError):
        calculate(inv([{"entity": "E", "category": "s1_stationary", "method": "factor", "factor_id": GAS,
                        "quantity": 10, "unit": "litres"}]))


def test_market_based_hierarchy():
    rec = {"entity": "E", "category": "s2_electricity", "method": "electricity", "kwh": 1_000_000,
           "instruments": [{"type": "REGO", "mwh": 400, "g_per_kwh": 0, "certificate": "c1"}], "evidence": "m"}
    r = calculate(inv([rec]))
    loc = F.electricity_location("GBR")["g_per_kwh"]
    rm = F.electricity_residual("GBR")["g_per_kwh"]
    assert r["totals_t"]["scope2_location"] == pytest.approx(1000 * loc / 1000, rel=1e-6)
    assert r["totals_t"]["scope2_market"] == pytest.approx(600_000 * rm / 1e6, rel=1e-6)
    # supplier-specific factor beats the residual mix
    r2 = calculate(inv([dict(rec, instruments=[], supplier_g_per_kwh=100)]))
    assert r2["totals_t"]["scope2_market"] == pytest.approx(100.0)


def test_no_residual_mix_falls_back_to_location_and_warns():
    r = calculate(inv([{"entity": "E", "category": "s2_electricity", "method": "electricity", "kwh": 1e6, "evidence": "m"}],
                      entities=[{"id": "E", "name": "E", "country": "IND"}]))
    assert r["totals_t"]["scope2_market"] == pytest.approx(r["totals_t"]["scope2_location"])
    assert any("no residual mix" in c["message"] for c in r["checks"])
    assert r["scope3_by_category_t"]["Fuel- and energy-related activities"] > 0   # T&D + upstream


def test_consolidation_approaches():
    ents = [{"id": "J", "name": "JV", "country": "GBR", "equity_pct": 40, "operational_control": False,
             "financial_control": False}]
    rec = [{"entity": "J", "category": "s1_stationary", "method": "factor", "factor_id": GAS, "quantity": 1e6, "evidence": "x"}]
    eq = calculate(inv(rec, ents, consolidation="equity_share"))["totals_t"]["scope1"]
    oc = calculate(inv(rec, ents, consolidation="operational_control"))["totals_t"]["scope1"]
    full = calculate(inv([dict(rec[0], entity="E")]))["totals_t"]["scope1"]
    assert eq == pytest.approx(0.4 * full, abs=0.002) and oc == 0


def test_completeness_requires_reason():
    d = {"organisation": "T", "reporting_year": 2025, "entities": [{"id": "E", "name": "E", "country": "GBR"}],
         "records": [{"entity": "E", "category": "s1_stationary", "method": "factor", "factor_id": GAS, "quantity": 1, "evidence": "x"}]}
    r = calculate(from_dict(d))
    assert sum(1 for c in r["checks"] if c["check"] == "completeness") == 14   # cat 3 reported via WTT
    assert any(c["level"] == "error" for c in r["checks"])


def test_direct_and_spend_records_and_evidence_warning():
    r = calculate(inv([
        {"entity": "E", "category": "s3_c4", "method": "direct", "kgco2e": 5000, "factor_source": "carrier (GLEC)"},
        {"entity": "E", "category": "s3_c1", "method": "spend", "naics": "541511", "amount": 100000, "currency": "GBP",
         "fx_to_usd": 1.27, "spend_year": 2025, "evidence": "AP"}]))
    assert r["scope3_by_category_t"]["Upstream transportation and distribution"] == pytest.approx(5.0)
    assert any("no evidence" in c["message"] for c in r["checks"])
    assert any("US economy" in c["message"] for c in r["checks"])
    with pytest.raises(InventoryError):
        calculate(inv([{"entity": "E", "category": "s3_c1", "method": "spend", "naics": "541511", "amount": 1,
                        "currency": "EUR"}]))


def test_base_year_recalculation():
    assert base_year_recalculation(10000, [{"type": "acquisition", "base_year_effect_t": 300}])["recalculation_required"] is False
    r = base_year_recalculation(10000, [{"type": "acquisition", "base_year_effect_t": 800}])
    assert r["recalculation_required"] and r["recalculated_base_year_t"] == 10800


def test_template_roundtrip_and_assurance_pack():
    d = parse_workbook(template_workbook(), "Example", 2025)
    assert {r["method"] for r in d["records"]} == {"factor", "electricity", "spend", "direct"}
    d["category_status"].update({k: v for k, v in ALL_EXCLUDED.items() if k not in d["category_status"]})
    i = from_dict(d)
    res = calculate(i)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(assurance_workbook(res, i)))
    assert {"Audit trail", "Factors used", "Completeness", "Checks", "Methodology"} <= set(wb.sheetnames)
    assert wb["Audit trail"].max_row == len(res["lines"]) + 1


# ── API ──────────────────────────────────────────────────────────────────────

def test_api_endpoints():
    assert c.get("/carbon/factors", params={"q": "natural gas"}).json()["results"]
    assert c.get("/carbon/grid-factor", params={"country": "USA", "us_state": "CA"}).json()["location"]["subregion"] == "CAMX"
    assert c.get("/carbon/template.xlsx").status_code == 200
    payload = {"organisation": "T", "reporting_year": 2025, "entities": [{"id": "E", "name": "E", "country": "GBR"}],
               "records": [{"entity": "E", "category": "s1_stationary", "method": "factor", "factor_id": GAS,
                            "quantity": 1000, "evidence": "x"}], "category_status": ALL_EXCLUDED}
    assert c.post("/carbon/inventory", json=payload).json()["totals_t"]["scope1"] > 0
    assert c.post("/carbon/inventory/assurance.xlsx", json=payload).status_code == 200
    assert c.post("/carbon/inventory/report", json=payload).json()["report"]["a_reporting_organisation"] == "T"
    bad = dict(payload, records=[{"entity": "X", "category": "s1_stationary", "method": "factor", "factor_id": GAS, "quantity": 1}])
    assert c.post("/carbon/inventory", json=bad).status_code == 422
    files = {"file": ("t.xlsx", template_workbook())}
    r = c.post("/carbon/inventory/upload", files=files, data={"organisation": "X", "reporting_year": "2025"})
    assert r.status_code == 200 and r.json()["result"]["totals_t"]["scope1"] > 0
    assert c.get("/carbon/dashboard").status_code == 200
