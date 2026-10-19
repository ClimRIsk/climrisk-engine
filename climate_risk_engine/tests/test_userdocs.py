"""The user's own documents: read, index, search, cite, forget — all local, none of it mixed into the public knowledge index."""
import os

import pytest

from cri.resident import semantic as S
from cri.resident import userdocs as UD

REPORT = (
    "Sustainability report 2025.\n\nThe Rotterdam terminal handled 4.2 million tonnes of bulk chemicals in 2025. Flooding risk at the terminal is managed by a quay raised to 3.5 metres above datum "
    "and by a pumping station that can remove 120 cubic metres per minute.\n\nScope 1 and 2 emissions fell 11 percent against the 2022 baseline, driven by the switch of the boiler house from fuel oil to "
    "biomethane. Scope 3 emissions are not yet reported.\n\nThe board approved a capital plan of 85 million euros for the Antwerp site, of which 22 million euros is for a new stormwater retention basin."
)


@pytest.fixture()
def home(tmp_path, monkeypatch):
    monkeypatch.setenv("CRI_HOME", str(tmp_path))
    monkeypatch.setenv("CRI_RAG_SEMANTIC", "0")                  # tests must not need the embedding model
    S._MODEL.update(m=None, tried=False)
    UD._CACHE["stamp"] = None
    yield tmp_path
    UD._CACHE["stamp"] = None


def test_txt_add_search_cite_and_forget(home):
    f = home / "report.txt"
    f.write_text(REPORT)
    r = UD.add(str(f))
    assert r["ok"] and r["passages"] >= 1
    assert [d["name"] for d in UD.list_docs()] == ["report.txt"]
    ps = UD.search("How high is the quay at the Rotterdam terminal?")
    assert ps and "3.5 metres" in ps[0]["text"] and ps[0]["doc"] == "report.txt"
    a = UD.ask("How much did Scope 1 and 2 emissions fall?")
    assert a["ok"] and "11 percent" in a["reply"] and "[1]" in a["reply"]
    assert UD.remove("report.txt") >= 1
    assert UD.search("quay") == [] and not UD.has_docs()


def test_not_covered_says_so(home):
    UD.add_text("note.txt", REPORT)
    a = UD.ask("What was the acquisition price of the Singapore refinery?")
    assert not a["ok"] and "not seem to cover" in a["reply"]


def test_docx_xlsx_csv(home):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_heading("Flood plan", 1)
    d.add_paragraph("The pumping station at the Antwerp site can remove 120 cubic metres per minute and is tested every quarter by the operations team on site.")
    d.save(home / "plan.docx")
    assert UD.add(str(home / "plan.docx"))["ok"]
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Assets"
    ws.append(["Site", "Elevation m", "Insured value"]); ws.append(["Rotterdam terminal", 1.8, "42 million euros"]); ws.append(["Antwerp site", 4.1, "30 million euros"])
    wb.save(home / "assets.xlsx")
    assert UD.add(str(home / "assets.xlsx"))["ok"]
    (home / "t.csv").write_text("site,note\nGdansk,quay flooded in 2023 for six hours during a storm surge event\n")
    assert UD.add(str(home / "t.csv"))["ok"]
    top = UD.search("insured value of the Rotterdam terminal")[0]
    assert top["doc"] == "assets.xlsx" and "42 million" in top["text"]
    assert "plan.docx" in {d["name"] for d in UD.list_docs()}


def test_bad_files_fail_plainly(home):
    assert not UD.add(str(home / "missing.pdf"))["ok"]
    (home / "x.bin").write_bytes(b"\x00\x01")
    r = UD.add(str(home / "x.bin"))
    assert not r["ok"] and "cannot read" in r["error"]


def test_documents_stay_out_of_the_public_index(home):
    UD.add_text("private.txt", REPORT)
    from cri.resident import knowledge as K
    assert not any("Rotterdam terminal" in p["text"] for p in K.search("Rotterdam terminal quay"))


def test_readding_replaces(home):
    UD.add_text("a.txt", REPORT)
    UD.add_text("a.txt", REPORT.replace("4.2 million", "9.9 million"))
    assert len(UD.list_docs()) == 1
    assert "9.9 million" in UD.search("tonnes of bulk chemicals handled")[0]["text"]


@pytest.mark.skipif(not os.environ.get("CRI_EMBED_DIR"), reason="needs the local embedding model (set CRI_EMBED_DIR)")
def test_semantic_finds_paraphrase(home, monkeypatch):
    monkeypatch.delenv("CRI_RAG_SEMANTIC", raising=False)
    S._MODEL.update(m=None, tried=False)
    UD.add_text("r.txt", REPORT)
    ps = UD.search("what is being spent on handling heavy rain water at the Belgian plant")
    assert ps and "stormwater retention basin" in ps[0]["text"]


def _wait(A, r, tries=100):
    import time
    jid = r.get("job")
    for _ in range(tries):
        st = A.JOBS[jid]
        if st["status"] != "running":
            return st["result"]
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def test_chat_add_ask_list_forget(home):
    from cri.resident import agent as A
    ag = A.get_agent()
    f = home / "annual.txt"
    f.write_text(REPORT)
    r = ag.handle("s1", f"add document {f}")
    assert r["intent"] == "docs" and "annual.txt" in _wait(A, r)["reply"]
    r = ag.handle("s1", "list my documents")
    assert "annual.txt" in str(r["blocks"])
    r = ag.handle("s1", "What does my report say about the capital plan for the Antwerp site?")
    out = _wait(A, r)
    assert "85 million" in out["reply"] and "[1]" in out["reply"]
    r = ag.handle("s1", "forget document annual.txt")
    assert "Removed" in r["reply"]


def test_chat_ignores_document_words_when_nothing_added(home):
    from cri.resident import agent as A
    r = A.get_agent().handle("s2", "what does my report say about flooding?")
    assert r.get("intent") != "docs"


def test_api_upload_roundtrip(home):
    import base64
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient
    from cri.api.main import app
    c = TestClient(app)
    r = c.post("/resident/docs", json={"name": "plan.txt", "data_b64": base64.b64encode(REPORT.encode()).decode()}).json()
    assert r["ok"] and r["passages"] >= 1
    assert c.post("/resident/docs", json={"name": "bad.bin", "data_b64": base64.b64encode(b"\x00").decode()}).json()["ok"] is False
    assert [d["name"] for d in c.get("/resident/docs").json()["documents"]] == ["plan.txt"]
    a = c.post("/resident/docs/ask", json={"question": "How high is the quay at the Rotterdam terminal?"}).json()
    assert a["ok"] and "3.5 metres" in a["reply"]
    assert c.delete("/resident/docs/plan.txt").json()["removed_passages"] >= 1
