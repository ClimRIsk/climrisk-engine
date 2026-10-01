"""
Assurance pack and data intake for the GHG inventory.

  assurance_workbook(result, inv) → .xlsx bytes: Summary, Scope 1, Scope 2 (location + market), Scope 3,
      Audit trail (every line: input → factor → working → result → evidence → data quality), Factors used,
      Completeness, Checks, Methodology. Written so an assurance provider can re-perform every number.
  report_sections(result, inv) → ISO 14064-1:2018 clause 9 report content (JSON), for the GHG statement.
  template_workbook() / parse_workbook(bytes) → client intake template (Entities, Activities, Electricity,
      Spend, Exclusions) with a legend and one example row per sheet.
"""
from __future__ import annotations

import io
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .inventory import CATEGORIES, Inventory, S3_CATEGORIES

HDR = Font(name="Arial", bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3A5F")
BODY = Font(name="Arial", size=10)
INPUT = Font(name="Arial", size=10, color="0000FF")


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR, HFILL, Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
            if isinstance(c.value, float):
                c.number_format = "#,##0.000"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def assurance_workbook(res: dict, inv: Inventory) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    T = res["totals_t"]
    rows = [
        ("Organisation", inv.organisation), ("Reporting year", inv.reporting_year),
        ("Consolidation approach", inv.consolidation.replace("_", " ")), ("GWP basis", inv.gwp),
        ("Scope 1 (t CO2e)", T["scope1"]), ("Scope 2 location-based (t CO2e)", T["scope2_location"]),
        ("Scope 2 market-based (t CO2e)", T["scope2_market"]), ("Scope 3 (t CO2e)", T["scope3"]),
        ("Total, location-based (t CO2e)", T["total_location"]), ("Total, market-based (t CO2e)", T["total_market"]),
        ("Data quality (weighted, 1 best – 5)", res["data_quality"]["weighted_score"]),
        ("Inventory fingerprint", res["fingerprint"]), ("Factor library version", res["factor_library"]),
        ("Standards", "; ".join(res["standards"])),
    ]
    ws.append(["Item", "Value"])
    for c in ws[1]:
        c.font, c.fill = HDR, HFILL
    for r in rows:
        ws.append(list(r))
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 40, 60
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
            if isinstance(c.value, float):
                c.number_format = "#,##0.000"

    lines = res["lines"]
    trail = [[l["record"], l.get("derived_from") or "", l["entity_name"], f"Scope {l['scope']}",
              CATEGORIES[l["category"]][1], l.get("description") or "", l["method"], l["working"],
              l["factor"].get("id"), l["factor"].get("source"), l["factor"].get("gwp", ""), l["kgco2e_gross"] / 1000,
              l["share"], l["t_co2e"], l["t_co2e_market"] if l["t_co2e_market"] is not None else "",
              l.get("evidence") or "", l["dq"], "; ".join(l["warnings"])] for l in lines]
    _sheet(wb, "Audit trail", ["Record", "Derived from", "Entity", "Scope", "Category", "Description", "Method",
                               "Working", "Factor ID", "Factor source", "GWP", "Gross t CO2e", "Consolidation share",
                               "Reported t CO2e (location)", "Reported t CO2e (market)", "Evidence", "DQ (1–5)", "Warnings"],
           trail, [8, 8, 22, 9, 30, 24, 12, 60, 34, 44, 18, 13, 11, 14, 14, 30, 8, 40])
    for sc, title in ((1, "Scope 1"), (3, "Scope 3")):
        _sheet(wb, title, ["Category", "Entity", "Record", "t CO2e"],
               [[CATEGORIES[l["category"]][1], l["entity_name"], l["record"], l["t_co2e"]] for l in lines if l["scope"] == sc],
               [44, 26, 10, 14])
    _sheet(wb, "Scope 2", ["Entity", "Record", "Location-based t CO2e", "Market-based t CO2e", "Working"],
           [[l["entity_name"], l["record"], l["t_co2e"], l["t_co2e_market"], l["working"]] for l in lines if l["scope"] == 2],
           [26, 10, 18, 18, 90])
    seen = {}
    for l in lines:
        f = l["factor"]
        seen.setdefault((f.get("id"), f.get("source")), [f.get("id"), f.get("label", ""), f.get("value"), f.get("unit", ""),
                                                          f.get("source"), f.get("gwp", "")])
    _sheet(wb, "Factors used", ["Factor ID", "Label", "Value", "Unit", "Source", "GWP"], list(seen.values()),
           [36, 50, 14, 22, 60, 24])
    _sheet(wb, "Completeness", ["Category", "Scope", "Status", "t CO2e", "Reason"],
           [[c["name"], c["scope"], c["status"], c["t_co2e"], c.get("reason") or ""] for c in res["completeness"]],
           [50, 8, 14, 14, 60])
    _sheet(wb, "Checks", ["Level", "Check", "Message"], [[c["level"], c["check"], c["message"]] for c in res["checks"]],
           [10, 14, 110])
    meth = [
        ["Boundary", f"{inv.consolidation.replace('_', ' ')}; entities and shares on the Audit trail sheet."],
        ["Scope 1", "Activity data × UK DESNZ 2026 factors (per-gas CH4/N2O restated to AR6 where available); "
                    "refrigerants at IPCC AR6 GWP100."],
        ["Scope 2 location-based", "UK: DESNZ 2026; US: EPA eGRID2023 subregion (CO2, CH4, N2O at AR6); Europe: AIB 2025 "
                                   "production mix; elsewhere Ember-derived direct factor unless a national factor is given."],
        ["Scope 2 market-based", "Contractual instruments → supplier-specific factor → residual mix (AIB 2025 Europe; "
                                 "Green-e 2025 US) → location-based where no residual mix exists (flagged)."],
        ["Scope 3", "Activity-based DESNZ factors, supplier-specific data, or spend-based USEEIO v1.3 (deflated to 2022 USD "
                    "with CPI-U). Category 3 derived from Scope 1 fuels (WTT) and Scope 2 electricity (T&D losses, upstream)."],
        ["Exclusions", "Every Scope 3 category is reported, excluded with a reason, or marked not relevant (Completeness)."],
        ["Recalculation", "Base year recalculated when structural/methodology changes exceed 5 % (GHG Protocol ch.5)."],
        ["Limitations", "Spend-based factors reflect the US economy; Ember-derived grid factors are estimates."],
    ]
    _sheet(wb, "Methodology", ["Topic", "Method"], meth, [26, 140])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_sections(res: dict, inv: Inventory) -> dict:
    """ISO 14064-1:2018 clause 9.3 GHG report content."""
    T = res["totals_t"]
    excluded = [c for c in res["completeness"] if c["status"] in ("excluded", "not_relevant")]
    return {
        "a_reporting_organisation": inv.organisation,
        "b_responsible_person": None,
        "c_reporting_period": str(inv.reporting_year),
        "d_organisational_boundaries": f"{inv.consolidation.replace('_', ' ')} approach; "
                                       f"{len(inv.entities)} entities (see Audit trail).",
        "e_reporting_boundaries": "Direct (Scope 1), energy indirect (Scope 2, location- and market-based), and other "
                                  "indirect (Scope 3) emissions as listed under category totals.",
        "f_direct_emissions_t": T["scope1"],
        "g_biogenic_co2": "Not reported separately unless supplied as a record (bioenergy CO2 is outside the scopes).",
        "h_removals": "None reported.",
        "i_excluded_sources": [{"category": c["name"], "status": c["status"], "reason": c["reason"]} for c in excluded],
        "j_indirect_emissions_t": {"scope2_location": T["scope2_location"], "scope2_market": T["scope2_market"],
                                   "scope3": T["scope3"], "scope3_by_category": res["scope3_by_category_t"]},
        "k_base_year": inv.base_year,
        "l_changes_or_recalculations": None,
        "m_quantification_methodologies": "Calculation-based: activity data × published emission factors (see Factors used).",
        "n_changes_in_methodologies": None,
        "o_factors_used": "UK DESNZ 2026; US EPA eGRID2023; Green-e 2025; AIB 2025; Ember; World Bank WDI; USEEIO v1.3.",
        "p_uncertainty": f"Weighted data-quality score {res['data_quality']['weighted_score']} "
                         f"(1 = metered primary data … 5 = proxy); primary-data share {res['data_quality']['share_primary']}.",
        "q_statement": "This GHG report has been prepared in accordance with ISO 14064-1:2018 and the GHG Protocol; "
                       "it has not been verified unless a verification statement is attached.",
        "r_verification": None,
        "gwp": inv.gwp,
    }


# ── intake template ──────────────────────────────────────────────────────────

TEMPLATE = {
    "Entities": (["id", "name", "country (ISO2/ISO3)", "equity_pct", "operational_control (Y/N)", "financial_control (Y/N)"],
                 ["UK1", "Example UK Ltd", "GB", 100, "Y", "Y"]),
    "Activities": (["id", "entity", "category", "factor_id", "quantity", "unit", "description", "evidence", "dq (1-5)"],
                   ["a1", "UK1", "s1_stationary", "desnz26:fuels-gaseous-fuels-natural-gas-kwh-gross-cv", 1200000,
                    "kWh (Gross CV)", "Head office boilers", "Gas invoices 2025", 1]),
    "Electricity": (["id", "entity", "kwh", "country", "us_state", "egrid_subregion", "instrument_type", "instrument_mwh",
                     "instrument_g_per_kwh", "certificate", "supplier_g_per_kwh", "evidence", "dq (1-5)"],
                    ["e1", "UK1", 850000, "GB", "", "", "REGO", 400, 0, "REGO-2025-001", "", "Meter readings", 1]),
    "Spend": (["id", "entity", "category", "naics", "amount", "currency", "fx_to_usd", "spend_year", "description",
               "evidence"],
              ["s1", "UK1", "s3_c1", "541511", 250000, "GBP", 1.27, 2025, "Software services", "AP ledger 2025"]),
    "Direct": (["id", "entity", "category", "kgco2e", "factor_source", "description", "evidence", "dq (1-5)"],
               ["d1", "UK1", "s3_c4", 18500, "Carrier-reported (GLEC)", "Inbound freight", "Carrier report Q1-Q4", 2]),
    "Exclusions": (["category", "status (excluded/not_relevant)", "reason"],
                   ["s3_c11", "not_relevant", "Products consume no energy in use"]),
}


def template_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    notes = [
        "ClimRisk GHG inventory intake template",
        "Blue cells in row 2 of each sheet are an example — replace them with your data (one row per record).",
        "Categories: " + ", ".join(CATEGORIES),
        "Scope 3 categories: " + "; ".join(f"s3_c{k} = {v}" for k, v in S3_CATEGORIES.items()),
        "factor_id: search the factor library (GET /carbon/factors?q=natural gas) and copy the id; unit must match.",
        "Electricity: give kWh and country; US sites also state or eGRID subregion. Instruments (GO/REGO/REC/PPA) "
        "with MWh, emission rate and certificate reference.",
        "Spend: 6-digit 2017 NAICS code, amount, currency and the FX rate to USD for the spend year.",
        "Direct: emissions already calculated by a supplier, CEMS or another tool, with the source named.",
        "Exclusions: every Scope 3 category without data must be listed with a reason.",
        "dq: 1 metered/invoiced, 2 supplier/primary, 3 partial/extrapolated, 4 spend/average, 5 proxy.",
    ]
    for n in notes:
        ws.append([n])
    ws.column_dimensions["A"].width = 150
    for c in ws["A"]:
        c.font = BODY
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    for name, (hdr, ex) in TEMPLATE.items():
        s = wb.create_sheet(name)
        s.append(hdr)
        s.append(ex)
        for c in s[1]:
            c.font, c.fill = HDR, HFILL
        for c in s[2]:
            c.font = INPUT
        for i in range(1, len(hdr) + 1):
            s.column_dimensions[get_column_letter(i)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _yn(v) -> bool:
    return str(v).strip().upper() in ("Y", "YES", "TRUE", "1")


def parse_workbook(data: bytes, organisation: str, reporting_year: int, consolidation: str = "operational_control",
                   gwp: str = "AR6", revenue_musd: Optional[float] = None, employees: Optional[float] = None) -> dict:
    wb = load_workbook(io.BytesIO(data), data_only=True)

    notes: list = []

    def rows(name):
        if name not in wb.sheetnames:
            return []
        ws = wb[name]
        vals = list(ws.iter_rows(values_only=True))
        if not vals:
            return []
        hdr = [str(h).split(" (")[0].strip() if h else "" for h in vals[0]]
        ex = TEMPLATE.get(name, (None, None))[1]
        out = []
        for r in vals[1:]:
            if not any(v not in (None, "") for v in r):
                continue
            norm = lambda xs: [None if v == "" else v for v in xs]
            if ex and norm(r[:len(ex)]) == norm(ex):        # the template's example row left in place
                notes.append(f"{name}: the template's example row was left in the file and was ignored")
                continue
            out.append(dict(zip(hdr, r)))
        return out

    ents = [{"id": str(r["id"]), "name": str(r["name"]), "country": str(r["country"]),
             "equity_pct": float(r.get("equity_pct") or 100), "operational_control": _yn(r.get("operational_control", "Y")),
             "financial_control": _yn(r.get("financial_control", "Y"))} for r in rows("Entities")]
    recs = []
    for r in rows("Activities"):
        recs.append({"id": str(r["id"]), "entity": str(r["entity"]), "year": reporting_year, "category": r["category"],
                     "method": "factor", "factor_id": r["factor_id"], "quantity": float(r["quantity"]),
                     "unit": r.get("unit"), "description": r.get("description"), "evidence": r.get("evidence"),
                     "dq": int(r.get("dq") or 2)})
    elec: dict = {}
    for r in rows("Electricity"):
        rid = str(r["id"])
        rec = elec.setdefault(rid, {"id": rid, "entity": str(r["entity"]), "year": reporting_year,
                                    "category": "s2_electricity", "method": "electricity", "kwh": float(r["kwh"]),
                                    "country": r.get("country"), "us_state": r.get("us_state") or None,
                                    "egrid_subregion": r.get("egrid_subregion") or None, "instruments": [],
                                    "evidence": r.get("evidence"), "dq": int(r.get("dq") or 2)})
        if r.get("supplier_g_per_kwh") not in (None, ""):
            rec["supplier_g_per_kwh"] = float(r["supplier_g_per_kwh"])
        if r.get("instrument_mwh") not in (None, "", 0):
            rec["instruments"].append({"type": r.get("instrument_type") or "EAC", "mwh": float(r["instrument_mwh"]),
                                       "g_per_kwh": float(r.get("instrument_g_per_kwh") or 0),
                                       "certificate": r.get("certificate")})
    recs += list(elec.values())
    for r in rows("Spend"):
        recs.append({"id": str(r["id"]), "entity": str(r["entity"]), "year": reporting_year, "category": r["category"],
                     "method": "spend", "naics": str(r["naics"]), "amount": float(r["amount"]),
                     "currency": r.get("currency") or "USD",
                     # no silent 1:1 — a non-USD amount without an FX rate is rejected by the calculation
                     "fx_to_usd": float(r["fx_to_usd"]) if r.get("fx_to_usd") not in (None, "") else
                     (1.0 if str(r.get("currency") or "USD").upper() == "USD" else None),
                     "spend_year": int(r.get("spend_year") or reporting_year), "description": r.get("description"),
                     "evidence": r.get("evidence"), "dq": 4})
    for r in rows("Direct"):
        recs.append({"id": str(r["id"]), "entity": str(r["entity"]), "year": reporting_year, "category": r["category"],
                     "method": "direct", "kgco2e": float(r["kgco2e"]), "factor_source": r.get("factor_source"),
                     "description": r.get("description"), "evidence": r.get("evidence"), "dq": int(r.get("dq") or 2)})
    status = {r["category"]: {"status": r.get("status"), "reason": r.get("reason")} for r in rows("Exclusions")}
    return {"organisation": organisation, "reporting_year": reporting_year, "consolidation": consolidation, "gwp": gwp,
            "entities": ents, "records": recs, "category_status": status, "revenue_musd": revenue_musd,
            "employees": employees, "intake_notes": notes}
